#!/usr/bin/env python3
"""
Deep-index the two WSMK workbooks into data/index.json for the Crew Performance platform.

Usage:
  python3 scripts/build_index.py "<path to WSMK PSC Status & Schedule...xlsx>" "<path to WSMK KPI.xlsx>"

Sources
  PSC  : "Database" sheet  (one row per deficiency; grouped here into inspections by vessel+date)
  UA   : "Tech Data" sheet (Unscheduled stoppage (Operational Delay) rows, Off-Hire hours)
  LTIF : "Crew Data1" sheet (LTIF / LTSF events with rank + description)
  Fleet: vessel lists (for per-vessel averages)
"""
import sys, json, re, datetime, collections
import openpyxl
import warnings
warnings.filterwarnings("ignore")

PSC_XLSX = sys.argv[1]
KPI_XLSX = sys.argv[2]
OUT = sys.argv[3] if len(sys.argv) > 3 else "data/index.json"

# ---------- helpers ----------
def norm(name):
    """Normalise vessel name: upper-case, collapse spaces, strip M/V prefixes."""
    if name is None:
        return None
    s = str(name).strip()
    s = re.sub(r"^(M/V|MV|M\.V\.)\s+", "", s, flags=re.I)
    s = re.sub(r"\s+", " ", s).upper()
    return s or None

def d(v):
    if isinstance(v, datetime.datetime):
        return v.date().isoformat()
    if isinstance(v, datetime.date):
        return v.isoformat()
    return None

def hours(v):
    """Convert Off-Hire cell to hours (float)."""
    if v is None:
        return 0.0
    if isinstance(v, datetime.timedelta):
        return round(v.total_seconds() / 3600, 2)
    if isinstance(v, datetime.time):
        return round(v.hour + v.minute / 60 + v.second / 3600, 2)
    if isinstance(v, datetime.datetime):
        # Excel time > 24h sometimes surfaces as datetime 1899-12-31 + days
        base = datetime.datetime(1899, 12, 30)
        return round((v - base).total_seconds() / 3600, 2)
    if isinstance(v, (int, float)):
        return round(float(v) * 24, 2)  # Excel serial fraction of a day
    s = str(v).strip()
    m = re.match(r"(?:(\d+)\s*day[s]?,\s*)?(\d+):(\d+)(?::(\d+))?", s)
    if m:
        days = int(m.group(1) or 0)
        return round(days * 24 + int(m.group(2)) + int(m.group(3)) / 60 + int(m.group(4) or 0) / 3600, 2)
    try:
        return round(float(s), 2)
    except Exception:
        return 0.0

def clean(s, n=400):
    if s is None:
        return ""
    return re.sub(r"\s+", " ", str(s)).strip()[:n]

# ---------- PSC ----------
wb = openpyxl.load_workbook(PSC_XLSX, read_only=True, data_only=True)
ws = wb["Database"]
insp = collections.OrderedDict()
blank = 0
for i, r in enumerate(ws.iter_rows(values_only=True)):
    if i == 0:
        continue
    if r[1] is None and r[4] is None:
        blank += 1
        if blank > 200:
            break
        continue
    blank = 0
    vessel = norm(r[1])
    date = d(r[4])
    if not vessel or not date:
        continue
    key = (vessel, date)
    if key not in insp:
        insp[key] = {
            "vessel": vessel, "date": date, "owner": clean(r[2], 40),
            "port": clean(r[7], 60), "country": clean(r[8], 60), "mou": clean(r[9], 40),
            "flag": clean(r[11], 10), "deficiencies": [], "detention": False,
        }
    rec = insp[key]
    code = clean(r[10], 20)
    desc = clean(r[16])
    loc = clean(r[17], 60)
    item = clean(r[19] or r[18], 80)
    if code.lower().startswith("no def") or code.lower().startswith("total") or (not code and not desc):
        continue
    if "30" in re.findall(r"\d+", code) or "detention" in desc.lower():
        rec["detention"] = True
    rec["deficiencies"].append({"code": code, "desc": desc, "location": loc, "item": item})

psc = list(insp.values())
for p in psc:
    p["def_count"] = len(p["deficiencies"])
    p["code17"] = sum(1 for x in p["deficiencies"] if "17" in re.findall(r"\d+", x["code"]))

# current fleet from schedule sheet
fleet_now = set()
try:
    ws = wb["WSMK Vessel Schedule"]
    for i, r in enumerate(ws.iter_rows(values_only=True, max_row=3000)):
        if i < 2:
            continue
        if r[0] and isinstance(r[0], str) and r[0].strip() and r[0].strip().upper() != "VESSEL NAME":
            fleet_now.add(norm(r[0]))
except Exception:
    pass

# ---------- KPI: Unplanned Unavailability ----------
wb = openpyxl.load_workbook(KPI_XLSX, read_only=True, data_only=True)
ws = wb["Tech Data"]
ua = []
for i, r in enumerate(ws.iter_rows(values_only=True)):
    if i < 4 or not r[1] or not r[4]:
        continue
    ev = str(r[4])
    if "unscheduled stoppage" not in ev.lower():
        continue
    date = d(r[3])
    if not date:
        continue
    ua.append({
        "vessel": norm(r[1]), "date": date, "hours": hours(r[6]),
        "remark": clean(r[7]), "kpi": clean(r[5], 20), "docmap": clean(r[11], 40),
    })

# ---------- KPI: LTIF ----------
ws = wb["Crew Data1"]
lti = []
for i, r in enumerate(ws.iter_rows(values_only=True)):
    if i < 4 or not r[1] or not r[4]:
        continue
    ev = str(r[4])
    date = d(r[3])
    if not date:
        continue
    typ = "LTIF" if "LTIF" in ev else ("LTSF" if "LTSF" in ev else clean(ev, 20))
    lti.append({"vessel": norm(r[1]), "date": date, "type": typ, "rank": clean(r[6], 20),
                "desc": clean(r[7], 200), "kpi": clean(r[5], 20)})

# LTIF/LTSF source is Crew Data1 only (per user requirement)

# ---------- Fleet size per year ----------
vessels_by_year = collections.defaultdict(set)
for p in psc:
    vessels_by_year[int(p["date"][:4])].add(p["vessel"])
for u in ua:
    vessels_by_year[int(u["date"][:4])].add(u["vessel"])
for x in lti:
    vessels_by_year[int(x["date"][:4])].add(x["vessel"])
# HSEQ Data covers every recorded event -> good proxy of the active fleet
ws = wb["HSEQ Data"]
for i, r in enumerate(ws.iter_rows(values_only=True)):
    if i < 4 or not r[1] or not d(r[3]):
        continue
    vessels_by_year[int(d(r[3])[:4])].add(norm(r[1]))
this_year = datetime.date.today().year
fleet = {str(y): len(v) for y, v in sorted(vessels_by_year.items())}
if fleet_now:
    fleet[str(this_year)] = max(fleet.get(str(this_year), 0), len(fleet_now))

all_vessels = sorted(set(p["vessel"] for p in psc) | set(u["vessel"] for u in ua) | set(x["vessel"] for x in lti) | fleet_now)

out = {
    "generated": datetime.datetime.now().isoformat(timespec="seconds"),
    "sources": {"psc": PSC_XLSX.split("/")[-1], "kpi": KPI_XLSX.split("/")[-1]},
    "fleet_by_year": fleet,
    "fleet_now": sorted(fleet_now),
    "vessels": all_vessels,
    "psc": sorted(psc, key=lambda x: x["date"]),
    "ua": sorted(ua, key=lambda x: x["date"]),
    "lti": sorted(lti, key=lambda x: x["date"]),
}
import os
os.makedirs(os.path.dirname(OUT) or ".", exist_ok=True)
with open(OUT, "w", encoding="utf-8") as f:
    json.dump(out, f, ensure_ascii=False, separators=(",", ":"))
print(f"PSC inspections: {len(psc)}  (deficiency rows: {sum(p['def_count'] for p in psc)})")
print(f"Unscheduled stoppage events: {len(ua)}  total hours: {round(sum(u['hours'] for u in ua),1)}")
print(f"LTI/LTS events: {len(lti)}")
print(f"Fleet by year: {fleet}")
print(f"Vessels: {len(all_vessels)}  -> {OUT}")
